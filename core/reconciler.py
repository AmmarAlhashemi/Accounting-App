from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from itertools import zip_longest
import re
import fitz
from openpyxl import load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd


def perform_reconciliation(excel_file, pdf_file, output_file):
    # 1. Parse Excel
    xl_list = _parse_excel_records(excel_file)
    if xl_list is None:
        return False

    # 2. Parse PDF
    pd_list = _parse_pdf_records(pdf_file)
    if pd_list is None:
        return False

    # 3. Match & Export
    return _match_and_export(xl_list, pd_list, output_file)


def _parse_excel_records(excel_file):
    try:
        workbook = load_workbook(excel_file, data_only=True)
        sheet = workbook.active
        all_rows = list(sheet.iter_rows(min_col=1, max_col=9, values_only=True))

        rec_indices = []
        for idx in range(len(all_rows) - 1, -1, -1):
            raw_row = all_rows[idx]
            if len(raw_row) >= 9 and raw_row[8] is not None:
                if "rec" in str(raw_row[8]).strip().lower():
                    rec_indices.append(idx)
                    if len(rec_indices) == 2:
                        break

        rec_indices.reverse()
        if not rec_indices:
            return []

        target_rows = all_rows[rec_indices[0] : rec_indices[-1] + 1]
        records = []

        for raw_row in target_rows:
            row_ah = raw_row[:8]
            if not any(row_ah):
                continue
            row = ["" if cell is None else cell for cell in row_ah]
            if str(row[-1]).strip() != "":
                continue

            # Date format
            date_val = row[0]
            if isinstance(date_val, datetime):
                formatted_date = f"{date_val.day} {date_val.strftime('%b')}"
            elif isinstance(date_val, str) and "/" in date_val:
                try:
                    dt = datetime.strptime(date_val.strip(), "%d/%m/%Y")
                    formatted_date = f"{dt.day} {dt.strftime('%b')}"
                except ValueError:
                    formatted_date = str(date_val).strip()
            else:
                formatted_date = str(date_val).strip()

            def parse_num(v):
                try:
                    return float(v)
                except (ValueError, TypeError):
                    return 0.0

            col4_val = parse_num(row[4])
            if col4_val.is_integer():
                col4_val = int(col4_val)

            amt_col6 = parse_num(row[6])
            amt_col5 = parse_num(row[5])
            total_amount = int(
                Decimal(str(amt_col6 if amt_col6 != 0 else amt_col5)).quantize(
                    Decimal("1"), rounding=ROUND_HALF_UP
                )
            )

            if not total_amount:
                continue

            type_code = str(row[1]).strip().upper()
            desc_raw = str(row[2]).strip()
            local_currency_val = parse_num(row[3])

            if "+" in desc_raw and "/" in desc_raw:
                parts = desc_raw.split("/", 1)
                clean_desc = parts[0].strip()
                after_slash = parts[1].strip()
                country = after_slash.split()[0] if after_slash else ""
                if country == "KEN":
                    country = "KENYA"
            elif "FEE" in type_code or "FEE" in desc_raw.upper():
                clean_desc = "FEE / COMMISSION"
                country = "USA"
            elif (
                any(k in type_code for k in ["TRN", "TRF"])
                or "BOA DEPOSIT" in desc_raw.upper()
            ):
                clean_desc = "AMERICA TRAN"
                country = "USA"
            elif local_currency_val > 0:
                formatted_amt = (
                    f"{int(local_currency_val):,}" if local_currency_val > 0 else ""
                )
                clean_desc = (
                    f"{desc_raw} {formatted_amt}".strip()
                    if formatted_amt
                    else desc_raw
                )
                country = "ETHIOPIA"
            else:
                clean_desc = desc_raw
                country = ""

            records.append(
                {
                    "date": formatted_date,
                    "desc": clean_desc,
                    "rate": col4_val,
                    "country": country,
                    "amount": total_amount,
                    "key": (col4_val, country, total_amount),
                }
            )
        return records
    except Exception as e:
        print(f"Error reading excel: {e}")
        return None


def _parse_pdf_records(pdf_file):
    try:
        records = []
        with fitz.open(pdf_file) as doc:
            for page in doc:
                data = page.get_text("dict")
                for block in data["blocks"]:
                    if block["type"] != 0:
                        continue
                    row = []
                    for line in block["lines"]:
                        text = "".join(span["text"] for span in line["spans"]).strip()
                        if text:
                            row.append(text)
                    if not row or not re.compile(r"^\d{1,2}\s+[A-Za-z]+").match(
                        row[0]
                    ):
                        continue
                    if len(row) != 5:
                        continue

                    date = row[0]
                    description = row[1].strip()
                    country = row[2].strip().upper()
                    raw_amount = row[3].replace(",", "").strip()

                    try:
                        amount = int(float(raw_amount))
                    except (ValueError, TypeError):
                        continue
                    if amount == 0:
                        continue

                    rate = 0
                    desc_upper = description.upper()
                    if "BANK OF AMERICA" in desc_upper:
                        description = "AMERICA TRAN"
                    elif "COMMISSION" in desc_upper:
                        description = "FEE / COMMISSION"
                    elif country == "ETHIOPIA":
                        match = re.compile(
                            r"(.+?)\s+([\d,]+)\s+BIRR\s+([\d.]+)", re.IGNORECASE
                        ).search(description)
                        if match:
                            description = f"{match.group(1).strip()} {match.group(2)}"
                            parsed_rate = float(match.group(3))
                            rate = (
                                int(parsed_rate)
                                if parsed_rate.is_integer()
                                else parsed_rate
                            )

                    records.append(
                        {
                            "date": date,
                            "desc": description,
                            "rate": rate,
                            "country": country,
                            "amount": amount,
                            "key": (rate, country, amount),
                        }
                    )
        return records
    except Exception as e:
        print(f"Error reading pdf: {e}")
        return None


def _match_and_export(xl_list, pd_list, output_file):
    for rec in xl_list + pd_list:
        rec["rate"] = float(rec["rate"])
        rec["amount"] = float(rec["amount"])
        rec["key"] = (rec["rate"], rec["country"], rec["amount"])

    all_records = [
        {
            "XLS Date": l["date"] if l else "",
            "XLS Description": l["desc"] if l else "",
            "XLS Rate": l["rate"] if l else "",
            "XLS Country": l["country"] if l else "",
            "XLS Amount": l["amount"] if l else "",
            "PDF Date": r["date"] if r else "",
            "PDF Description": r["desc"] if r else "",
            "PDF Rate": r["rate"] if r else "",
            "PDF Country": r["country"] if r else "",
            "PDF Amount": r["amount"] if r else "",
        }
        for r, l in zip_longest(pd_list, xl_list, fillvalue=None)
    ]

    pd_remaining, xl_remaining = list(pd_list), list(xl_list)
    matched_pairs = []
    i = 0

    while i < len(pd_remaining):
        r_item = pd_remaining[i]
        match_found = False
        for j, l_item in enumerate(xl_remaining):
            if r_item["key"] == l_item["key"]:
                matched_pairs.append(
                    {
                        "XLS Date": l_item["date"],
                        "XLS Description": l_item["desc"],
                        "XLS Rate": l_item["rate"],
                        "XLS Country": l_item["country"],
                        "XLS Amount": l_item["amount"],
                        "PDF Date": r_item["date"],
                        "PDF Description": r_item["desc"],
                        "PDF Rate": r_item["rate"],
                        "PDF Country": r_item["country"],
                        "PDF Amount": r_item["amount"],
                    }
                )
                pd_remaining.pop(i)
                xl_remaining.pop(j)
                match_found = True
                break
        if not match_found:
            i += 1

    unmatched_rows = [
        {
            "XLS Date": (
                xl_remaining[k]["date"] if k < len(xl_remaining) else ""
            ),
            "XLS Description": (
                xl_remaining[k]["desc"] if k < len(xl_remaining) else ""
            ),
            "XLS Rate": (
                xl_remaining[k]["rate"] if k < len(xl_remaining) else ""
            ),
            "XLS Country": (
                xl_remaining[k]["country"] if k < len(xl_remaining) else ""
            ),
            "XLS Amount": (
                xl_remaining[k]["amount"] if k < len(xl_remaining) else ""
            ),
            "PDF Date": (
                pd_remaining[k]["date"] if k < len(pd_remaining) else ""
            ),
            "PDF Description": (
                pd_remaining[k]["desc"] if k < len(pd_remaining) else ""
            ),
            "PDF Rate": (
                pd_remaining[k]["rate"] if k < len(pd_remaining) else ""
            ),
            "PDF Country": (
                pd_remaining[k]["country"] if k < len(pd_remaining) else ""
            ),
            "PDF Amount": (
                pd_remaining[k]["amount"] if k < len(pd_remaining) else ""
            ),
        }
        for k in range(max(len(pd_remaining), len(xl_remaining)))
    ]

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        pd.DataFrame(unmatched_rows).to_excel(
            writer, sheet_name="Unmatched", index=False
        )
        pd.DataFrame(matched_pairs).to_excel(
            writer, sheet_name="Matched", index=False
        )
        pd.DataFrame(all_records).to_excel(
            writer, sheet_name="All Records", index=False
        )

        wb = writer.book
        header_fill = PatternFill(
            start_color="1F497B", end_color="1F497B", fill_type="solid"
        )
        header_font = Font(color="FFFFFF", bold=True)
        unmatched_fill = PatternFill(
            start_color="E6B8B7", end_color="E6B8B7", fill_type="solid"
        )
        matched_fill = PatternFill(
            start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"
        )
        thick_right_side = Side(style="thick", color="000000")

        for sheet_name in ["Unmatched", "Matched", "All Records"]:
            ws = wb[sheet_name]
            max_r = ws.max_row

            for col in range(1, 11):
                cell = ws.cell(row=1, column=col)
                cell.font, cell.fill = header_font, header_fill

            for r in range(1, max_r + 1):
                c = ws.cell(row=r, column=5)
                c.border = Border(
                    left=c.border.left,
                    right=thick_right_side,
                    top=c.border.top,
                    bottom=c.border.bottom,
                )

            if sheet_name == "Unmatched":
                for r in range(2, max_r + 1):
                    for col in range(1, 11):
                        cell = ws.cell(row=r, column=col)
                        cell.fill = unmatched_fill
            elif sheet_name == "Matched":
                for r in range(2, max_r + 1):
                    for col in range(1, 11):
                        cell = ws.cell(row=r, column=col)
                        cell.fill = matched_fill

            for col in range(1, 11):
                max_len = max(
                    (
                        len(str(ws.cell(row=r, column=col).value or ""))
                        for r in range(1, ws.max_row + 1)
                    ),
                    default=10,
                )
                ws.column_dimensions[get_column_letter(col)].width = max(
                    max_len + 3, 12
                )

        ws_all = wb["All Records"]
        last_row = len(all_records) + 1
        for country, offset in [("KENYA", 4), ("ETHIOPIA", 5), ("USA", 6)]:
            r = last_row + offset
            for col_idx in [4, 9]:
                c = ws_all.cell(row=r, column=col_idx, value=country)
                c.font = Font(color="FFFFFF", bold=True)
                c.fill = PatternFill(
                    start_color="366092", end_color="366092", fill_type="solid"
                )

            c_e = ws_all.cell(
                row=r,
                column=5,
                value=f'=SUMIF(D2:D{last_row}, "{country}", E2:E{last_row})',
            )
            c_e.font = Font(color="000000", bold=True)
            c_e.fill = PatternFill(
                start_color="B8CCE4", end_color="B8CCE4", fill_type="solid"
            )

            c_j = ws_all.cell(
                row=r,
                column=10,
                value=f'=SUMIF(I2:I{last_row}, "{country}", J2:J{last_row})',
            )
            c_j.font = Font(color="000000", bold=True)
            c_j.fill = PatternFill(
                start_color="B8CCE4", end_color="B8CCE4", fill_type="solid"
            )

    return True