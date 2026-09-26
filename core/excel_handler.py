from copy import copy
import os
import re
import openpyxl


class ExcelHandler:

    def __init__(self, file_path):
        self.file_path = file_path

    def load_rate_and_ac_no(self):
        """Extract last rate and AC number from workbook."""
        rate, ac_no = None, None
        if not os.path.exists(self.file_path):
            return rate, ac_no

        wb = None
        try:
            wb = openpyxl.load_workbook(self.file_path, data_only=True)
            ws = wb.active
            rate_found, ac_found = False, False

            for i in range(ws.max_row, 1, -1):
                if not rate_found:
                    val = ws.cell(row=i, column=5).value
                    if isinstance(val, (int, float)):
                        rate = str(val)
                        rate_found = True

                if not ac_found:
                    text = ws.cell(row=i, column=3).value
                    if isinstance(text, str) and text.startswith("BOA"):
                        m = re.search(
                            r"A/C\s*NO\.?\s*(\d+)", text, re.IGNORECASE
                        )
                        if m:
                            ac_no = m.group(1)
                            ac_found = True

                if rate_found and ac_found:
                    break
        except Exception:
            pass
        finally:
            if wb:
                wb.close()
        return rate, ac_no

    def process_excel_batch(self, transaction_cache, is_new_balance):
        """Insert batch transactions into Excel workbook."""
        wb = openpyxl.load_workbook(self.file_path)
        ws = wb.active

        if not is_new_balance:
            for entry in transaction_cache:
                self._insert_entry(ws, entry)
            self._update_totals_formulas(ws)
        else:
            old_totals_row = ws.max_row
            style_source_row = old_totals_row - 1

            top_blank_row = old_totals_row + 1
            bottom_blank_row = old_totals_row + 2
            new_totals_row = old_totals_row + 3

            ws.cell(row=top_blank_row, column=1, value="")
            ws.cell(row=bottom_blank_row, column=1, value="")

            self._copy_row_style(ws, style_source_row, top_blank_row)
            self._copy_row_style(ws, style_source_row, bottom_blank_row)
            self._copy_row_style(
                ws, old_totals_row, new_totals_row, copy_values=True
            )

            first_data_row = top_blank_row + 1

            for entry in transaction_cache:
                self._insert_entry(ws, entry)

            last_data_row = ws.max_row - 2
            final_totals_row = ws.max_row

            ws.cell(
                row=final_totals_row, column=4
            ).value = f"=SUM(D{first_data_row}:D{last_data_row})"
            ws.cell(
                row=final_totals_row, column=6
            ).value = f"=SUM(F{first_data_row}:F{last_data_row})"
            ws.cell(
                row=final_totals_row, column=7
            ).value = f"=SUM(G{first_data_row}:G{last_data_row})"
            ws.cell(
                row=final_totals_row, column=8
            ).value = f"=SUM(H{old_totals_row}+F{final_totals_row}-G{final_totals_row})"

        wb.save(self.file_path)

    def _insert_entry(self, ws, entry):
        target_row = ws.max_row - 1
        if entry["type"] == "TRN":
            ws.insert_rows(target_row)
            self._fill_row(ws, target_row, entry, is_fee=False)
            self._apply_row_style(ws, target_row)

            target_row += 1
            ws.insert_rows(target_row)
            self._fill_row(ws, target_row, entry, is_fee=True)
            self._apply_row_style(ws, target_row)
        else:
            ws.insert_rows(target_row)
            self._fill_row(ws, target_row, entry, is_fee=False)
            self._apply_row_style(ws, target_row)

    def _fill_row(self, ws, row, data, is_fee):
        ws.cell(row=row, column=1, value=data["date"])
        if data["type"] == "TRN":
            if not is_fee:
                ws.cell(
                    row=row,
                    column=2,
                    value=(
                        f"TRN.{data['seq']}"
                        if str(data["seq"]).isdigit()
                        else data["seq"]
                    ),
                )
                ws.cell(row=row, column=3, value=data["desc"])
                ws.cell(row=row, column=4, value=0)
                ws.cell(row=row, column=6, value=data["amount"])
                ws.cell(row=row, column=7, value=0)
            else:
                ws.cell(
                    row=row,
                    column=2,
                    value=(
                        "TRN.FEE"
                        if str(data["seq"]).isdigit()
                        else f"{data['seq']}.FEE"
                    ),
                )
                ws.cell(
                    row=row,
                    column=3,
                    value=f"DEPOSIT FEE 2% ON ${data['amount']:,.0f}",
                )
                ws.cell(row=row, column=4, value=0)
                ws.cell(row=row, column=6, value=0)
                ws.cell(row=row, column=7, value=f"=F{row-1}*2%")
        elif data["type"] == "OTHERS":
            ws.cell(row=row, column=2, value="ORDER")
            country = data["seq"].upper()
            ws.cell(
                row=row,
                column=3,
                value=f"{data['desc']} /{country} ${data['amount']:,.0f}+{data['amount']*0.02:.2f}",
            )
            ws.cell(row=row, column=4, value=0)
            ws.cell(row=row, column=6, value=0)
            ws.cell(
                row=row, column=7, value=data["amount"] + data["amount"] * 0.02
            )
        else:  # ORDER
            ws.cell(row=row, column=2, value=f"ORDER-{data['seq']}")
            ws.cell(row=row, column=3, value=data["desc"])
            ws.cell(row=row, column=4, value=data["amount"])
            ws.cell(row=row, column=5, value=data["rate"])
            ws.cell(row=row, column=6, value=0)
            ws.cell(row=row, column=7, value=f"=D{row}/E{row}")

    def _apply_row_style(self, ws, row):
        source_row = row - 1
        if source_row < 1:
            return
        ws.row_dimensions[row].height = 15.75
        for col in range(1, 9):
            source_cell = ws.cell(row=source_row, column=col)
            target_cell = ws.cell(row=row, column=col)
            if source_cell.has_style:
                target_cell.font = copy(source_cell.font)
                target_cell.border = copy(source_cell.border)
                target_cell.fill = copy(source_cell.fill)
                target_cell.number_format = copy(source_cell.number_format)
                target_cell.alignment = copy(source_cell.alignment)

    def _copy_row_style(self, ws, source_row, target_row, copy_values=False):
        if source_row < 1:
            return
        ws.row_dimensions[target_row].height = (
            ws.row_dimensions[source_row].height or 15.75
        )
        for col in range(1, 9):
            source_cell = ws.cell(row=source_row, column=col)
            target_cell = ws.cell(row=target_row, column=col)
            if copy_values:
                target_cell.value = source_cell.value
            if source_cell.has_style:
                target_cell.font = copy(source_cell.font)
                target_cell.border = copy(source_cell.border)
                target_cell.fill = copy(source_cell.fill)
                target_cell.number_format = copy(source_cell.number_format)
                target_cell.alignment = copy(source_cell.alignment)

    def _update_totals_formulas(self, ws):
        last_row = ws.max_row
        for col, let in [(4, "D"), (6, "F"), (7, "G")]:
            cell_val = ws.cell(row=last_row, column=col).value
            if cell_val and ":" in cell_val:
                part1 = cell_val.split(":")[0]
                ws.cell(
                    row=last_row, column=col
                ).value = f"{part1}:{let}{last_row-1})"
        part1 = ws.cell(row=last_row, column=8).value.split("+")[0]
        ws.cell(
            row=last_row, column=8
        ).value = f"{part1}+F{last_row}-G{last_row})"